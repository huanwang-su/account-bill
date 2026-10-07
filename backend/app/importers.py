"""Extensible statement parsers. File content is data, never executable instructions."""

from abc import ABC, abstractmethod
from dataclasses import dataclass, asdict
from collections import defaultdict
import csv, io, json, re, hashlib
from .domain import DomainError, Money, Clock
from .models import (
    Account,
    Category,
    ImportBatch,
    ImportRow,
    ImportSource,
    Transaction,
    BillDisposition,
)
from .services import TransactionService, Repository
from .account_matching import AlipayAccountResolver, StatementAccountResolver
from .bill_status import BillStatusService


@dataclass
class ParsedStatement:
    owner: str
    rows: list
    metadata: dict


class StatementParser(ABC):
    @abstractmethod
    def parse(self, content: bytes, filename: str) -> ParsedStatement: ...


class AlipayParser(StatementParser):
    VERSION = "1.0"
    REQUIRED = {
        "交易时间",
        "交易分类",
        "交易对方",
        "商品说明",
        "收/支",
        "金额",
        "收/付款方式",
        "交易状态",
        "交易订单号",
    }
    ALIASES = {"金额(元)": "金额", "金额（元）": "金额", "交易号": "交易订单号"}

    def _table(self, content, filename):
        if filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook
            import zipfile

            try:
                with zipfile.ZipFile(io.BytesIO(content)) as z:
                    if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024:
                        raise DomainError("Excel 解压后过大")
                wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                try:
                    return [
                        list(r) for r in wb.active.iter_rows(values_only=True)
                    ], "xlsx"
                finally:
                    wb.close()
            except DomainError:
                raise
            except Exception:
                raise DomainError("无法读取 Excel 文件，请使用有效的 .xlsx 文件")
        if not filename.lower().endswith(".csv"):
            raise DomainError("目前支持支付宝 CSV 和相同表头的 XLSX")
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                text = content.decode(encoding)
                return list(csv.reader(io.StringIO(text))), encoding
            except UnicodeDecodeError:
                pass
        raise DomainError("无法识别账单编码，请使用支付宝导出的原文件")

    def parse(self, content, filename):
        table, encoding = self._table(content, filename)
        owner = ""
        header = None
        start = 0
        declared = {}
        declared_count = None
        for i, row in enumerate(table):
            clean = [str(v or "").strip().lstrip("\ufeff") for v in row]
            joined = ",".join(clean)
            if joined.startswith("支付宝账户："):
                owner = joined.split("：", 1)[1].strip()
            m = re.search(r"共\s*(\d+)\s*笔记录", joined)
            if m:
                declared_count = int(m[1])
            m = re.match(r"(收入|支出|不计收支)[：:]\s*(\d+)笔\s*([\d,.]+)元", joined)
            if m:
                declared[m[1]] = {"count": int(m[2]), "amount": Money.parse(m[3])}
            normalized = [self.ALIASES.get(v, v) for v in clean]
            if self.REQUIRED.issubset(set(normalized)):
                header = normalized
                start = i + 1
                break
        if header is None:
            raise DomainError("未找到支付宝交易表头，请上传完整的支付宝交易明细")
        rows = []
        totals = defaultdict(lambda: {"count": 0, "amount": 0})
        for i, row in enumerate(table[start:], start + 1):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            if (
                str(row[0] or "")
                .strip()
                .startswith(("-----", "合计", "总计", "特别提示"))
            ):
                break
            raw = {
                k: str(row[j] if j < len(row) and row[j] is not None else "").strip()
                for j, k in enumerate(header)
                if k
            }
            if raw.get("交易时间") == "交易时间":
                continue
            d = {
                "line": i,
                "raw": raw,
                "direction": raw.get("收/支", ""),
                "status": raw.get("交易状态", ""),
                "counterparty": raw.get("交易对方", ""),
                "note": raw.get("商品说明", ""),
                "method": raw.get("收/付款方式", ""),
                "source_category": raw.get("交易分类", ""),
                "order_id": raw.get("交易订单号", ""),
                "occurred_at": raw.get("交易时间", ""),
                "amount": 0,
                "error": "",
            }
            try:
                d["occurred_at"] = Clock.normalize(d["occurred_at"])
                d["amount"] = Money.parse(raw.get("金额", ""))
                if d["amount"] < 0:
                    raise DomainError("金额不能为负数，请人工检查收支方向")
                if d["direction"] not in ("收入", "支出", "不计收支"):
                    raise DomainError("未知收支方向")
                totals[d["direction"]]["count"] += 1
                totals[d["direction"]]["amount"] += d["amount"]
            except DomainError as e:
                d["error"] = str(e)
            event = (
                ("refund|" + d["occurred_at"] + "|" + str(d["amount"]))
                if "退款" in d["status"]
                else "payment"
            )
            identity = owner + "|" + d["order_id"] + "|" + event
            # Without a provider ID, deduplicate the exact same file only; leave
            # potentially identical cash transactions in other files for review.
            d["source_key"] = (
                "alipay:" + hashlib.sha256(identity.encode()).hexdigest()
                if d["order_id"]
                else "file:" + hashlib.sha256(content).hexdigest() + ":" + str(i)
            )
            rows.append(d)
        if not rows:
            raise DomainError("文件没有交易记录")
        warnings = []
        if declared_count is not None and declared_count != len(rows):
            warnings.append("声明笔数与解析笔数不一致，请检查原文件")
        for direction, expected in declared.items():
            if dict(totals[direction]) != expected:
                warnings.append(direction + "明细合计与文件摘要不同，请按原始明细核对")
        return ParsedStatement(
            owner,
            rows,
            {
                "parser": "alipay",
                "version": self.VERSION,
                "encoding": encoding,
                "declared_count": declared_count,
                "declared_totals": declared,
                "totals": dict(totals),
                "warnings": warnings,
            },
        )


class ParserRegistry:
    def __init__(self):
        from .statement_parsers import WechatParser, PouchParser, CmbParser, BocParser

        self.parsers = {
            "alipay": AlipayParser(),
            "wechat": WechatParser(),
            "alipay_pouch": PouchParser(),
            "cmb": CmbParser(),
            "boc": BocParser(),
        }

    def detect(self, content, filename):
        from .statement_parsers import SpreadsheetReader, BankPdfReader

        if content.startswith(b"%PDF-"):
            with BankPdfReader.open(content) as pdf:
                text = pdf.pages[0].extract_text() or ""
            if "招商银行交易流水" in text:
                return "cmb"
            if "中国银行交易流水明细清单" in text:
                return "boc"
        else:
            table = SpreadsheetReader.read(content, filename)
            text = "\n".join(
                " ".join(SpreadsheetReader.text(v) for v in r) for r in table[:40]
            )
            if "支付宝小荷包账户ID" in text:
                return "alipay_pouch"
            if "微信支付账单明细" in text:
                return "wechat"
            if any(
                AlipayParser.REQUIRED.issubset(
                    {
                        AlipayParser.ALIASES.get(
                            SpreadsheetReader.text(v), SpreadsheetReader.text(v)
                        )
                        for v in r
                    }
                )
                for r in table[:40]
            ):
                return "alipay"
        raise DomainError(
            "无法识别账单来源，支持支付宝、微信、小荷包 XLSX/CSV 及招商银行、中国银行原始 PDF"
        )

    def get(self, name):
        if name not in self.parsers:
            raise DomainError("不支持的账单来源")
        return self.parsers[name]


class ImportService:
    def __init__(self, database, registry=None):
        self.db = database
        self.registry = registry or ParserRegistry()
        self.account_resolver = AlipayAccountResolver()

    def upload(self, content, filename, provider="alipay"):
        digest = hashlib.sha256(content).hexdigest()
        provider = (
            self.registry.detect(content, filename) if provider == "auto" else provider
        )
        statement = self.registry.get(provider).parse(content, filename)
        with self.db.write() as s:
            if provider == "alipay":
                BillStatusService(self.db).apply_rows(s, statement.rows)
            old = s.query(ImportBatch).filter_by(digest=digest).first()
            if old:
                return {"id": old.id, "existing": True}
            batch = ImportBatch(
                filename=filename,
                digest=digest,
                created_at=Clock.now(),
                owner=statement.owner,
                metadata_json=json.dumps(statement.metadata, ensure_ascii=False),
            )
            s.add(batch)
            s.flush()
            cats = s.query(Category).filter_by(active=True).all()
            catmap = {
                "餐饮美食": "餐饮",
                "日用百货": "日用品",
                "交通出行": "交通",
                "服饰装扮": "服饰",
                "投资理财": "理财收益",
            }
            seen = set()
            for d in statement.rows:
                warning = []
                decision = {}
                state = "REVIEW"
                method = d["method"]
                match = StatementAccountResolver().resolve(s, provider, statement, d)
                if d.get("review_reason"):
                    warning.append(d["review_reason"])
                if match:
                    decision["account_id"] = match.id
                else:
                    warning.append("请选择实际收付款账户")
                kind = {"收入": "INCOME", "支出": "EXPENSE"}.get(d["direction"])
                if kind:
                    decision["kind"] = kind
                    category = next(
                        (
                            c
                            for c in cats
                            if c.kind == kind
                            and c.name
                            == catmap.get(d["source_category"], d["source_category"])
                        ),
                        None,
                    )
                    if not category:
                        category = next(
                            c for c in cats if c.name == "未知用途" and c.kind == kind
                        )
                    decision["category_id"] = category.id
                else:
                    warning.append("支付宝标为不计收支，请确认支出、收入、转账或忽略")
                if "&" in method or "因公付" in method:
                    warning.append("包含优惠、混合支付或企业付款，请核实个人实际金额")
                if d["status"] not in ("交易成功", "支付成功"):
                    warning.append(
                        "交易状态：" + d["status"] + "，请核实资金是否实际变动"
                    )
                if not d["order_id"]:
                    warning.append("缺少外部交易号，请检查是否重复")
                # Cross-file same time/amount/payee is a warning, never an automatic skip.
                candidates = s.query(ImportRow).all() if not d["order_id"] else []
                if any(
                    (
                        lambda x: x.get("occurred_at") == d["occurred_at"]
                        and x.get("amount") == d["amount"]
                        and x.get("counterparty") == d["counterparty"]
                    )(json.loads(r.data))
                    for r in candidates
                ):
                    warning.append("存在同时间、金额和对方的疑似重复记录")
                if d["error"]:
                    state = "ERROR"
                    warning = [d["error"]]
                elif BillStatusService.excluded(s, d["source_key"]):
                    state = BillStatusService.routing(s, d["source_key"])
                    warning = [s.get(BillDisposition, d["source_key"]).reason]
                elif d["amount"] == 0:
                    state = "IGNORED"
                    warning = ["零金额记录保留原始数据，不生成资金流水"]
                elif d["status"] == "交易关闭":
                    state = "IGNORED"
                    warning = ["交易关闭，默认忽略"]
                elif s.get(ImportSource, d["source_key"]) or d["source_key"] in seen:
                    state = "DUPLICATE"
                    warning = ["外部交易已导入或在文件中重复"]
                elif provider == "alipay" and d["direction"] == "不计收支":
                    state = "IGNORED"
                    warning = [BillStatusService.REASON]
                elif not warning:
                    state = "READY"
                seen.add(d["source_key"])
                s.add(
                    ImportRow(
                        batch_id=batch.id,
                        line=d["line"],
                        data=json.dumps(d, ensure_ascii=False),
                        source_key=d["source_key"],
                        status=state,
                        warning="；".join(warning),
                        decision=json.dumps(decision, ensure_ascii=False),
                    )
                )
            return {"id": batch.id, "existing": False}

    def link_existing_pouches(self):
        """Fill missing account associations without posting or overriding user decisions."""
        linked = 0
        with self.db.write() as s:
            for batch in s.query(ImportBatch):
                if json.loads(batch.metadata_json).get("parser") != "alipay":
                    continue
                for row in s.query(ImportRow).filter_by(
                    batch_id=batch.id, transaction_id=None
                ):
                    decision = json.loads(row.decision)
                    method = json.loads(row.data).get("method", "")
                    if decision.get("account_id") or not method.startswith(
                        "支付宝小荷包"
                    ):
                        continue
                    account = self.account_resolver.resolve(s, batch.owner, method)
                    if account:
                        decision["account_id"] = account.id
                        row.decision = json.dumps(decision, ensure_ascii=False)
                        row.warning = "；".join(
                            w
                            for w in row.warning.split("；")
                            if w != "请选择实际收付款账户"
                        )
                        linked += 1
        return linked

    def batches(self):
        with self.db.read() as s:
            result = []
            for b in s.query(ImportBatch).order_by(ImportBatch.id.desc()):
                counts = defaultdict(int)
                for r in s.query(ImportRow).filter_by(batch_id=b.id):
                    counts[r.status] += 1
                result.append(
                    dict(
                        id=b.id,
                        filename=b.filename,
                        created_at=b.created_at,
                        counts=dict(counts),
                        metadata=json.loads(b.metadata_json),
                    )
                )
            return result

    def rows(self, batch_id):
        with self.db.read() as s:
            Repository(s).require(ImportBatch, batch_id)
            return [
                dict(
                    **{
                        **json.loads(r.data),
                        "source_status": json.loads(r.data)["status"],
                        "status": r.status,
                    },
                    id=r.id,
                    batch_id=r.batch_id,
                    warning=r.warning,
                    decision=json.loads(r.decision),
                    transaction_id=r.transaction_id
                )
                for r in s.query(ImportRow)
                .filter_by(batch_id=batch_id)
                .order_by(ImportRow.line)
            ]

    def decide(self, row_id, data):
        with self.db.write() as s:
            r = Repository(s).require(ImportRow, row_id)
            if r.status in ("POSTED", "DUPLICATE", "ERROR", "REVERSED"):
                raise DomainError("该记录已入账、已撤销、重复或存在解析错误，不能修改")
            BillStatusService.require_normal(s, r.source_key)
            if data.get("ignore"):
                r.status = "IGNORED"
                return {"ok": True}
            d = json.loads(r.data)
            decision = {
                k: data[k]
                for k in (
                    "kind",
                    "account_id",
                    "to_account_id",
                    "category_id",
                    "amount",
                    "related_id",
                    "tags",
                )
                if k in data
            }
            payload = self._payload(r, d, decision)
            TransactionService(self.db)._postings(s, payload)
            r.decision = json.dumps(decision, ensure_ascii=False)
            r.status = "READY"
            r.warning = "已人工确认"
            return {"ok": True}

    def _payload(self, r, d, decision):
        return {
            **decision,
            "occurred_at": d["occurred_at"],
            "amount": decision.get(
                "amount",
                format(__import__("decimal").Decimal(d["amount"]) / 100, ".2f"),
            ),
            "counterparty": d["counterparty"],
            "note": d["note"],
            "source": r.source_key.split(":", 1)[0],
            "idempotency_key": r.source_key,
            "_bill_id": r.source_key,
        }

    def commit(self, batch_id):
        result = {"posted": 0, "duplicates": 0, "failed": []}
        with self.db.read() as s:
            Repository(s).require(ImportBatch, batch_id)
            ids = [
                r.id
                for r in s.query(ImportRow).filter_by(batch_id=batch_id, status="READY")
            ]
        for row_id in ids:
            try:
                with self.db.write() as s:
                    r = Repository(s).require(ImportRow, row_id)
                    if r.status != "READY":
                        continue
                    existing = s.get(ImportSource, r.source_key)
                    if existing:
                        r.status = "DUPLICATE"
                        r.transaction_id = existing.transaction_id
                        result["duplicates"] += 1
                        continue
                    tx = TransactionService(self.db).post_in(
                        s, self._payload(r, json.loads(r.data), json.loads(r.decision))
                    )
                    if not s.get(ImportSource, r.source_key):
                        s.add(
                            ImportSource(
                                key=r.source_key, transaction_id=tx.id, row_id=r.id
                            )
                        )
                    r.transaction_id = tx.id
                    r.status = "POSTED"
                    result["posted"] += 1
            except DomainError as e:
                result["failed"].append({"row_id": row_id, "error": str(e)})
        return result
