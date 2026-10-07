"""Source adapters for local exports. No remote services or executable file content."""

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from .domain import Clock, DomainError, Money
from .importers import ParsedStatement, StatementParser


class SpreadsheetReader:
    @staticmethod
    def read(content, filename):
        if filename.lower().endswith(".csv"):
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    return list(csv.reader(io.StringIO(content.decode(encoding))))
                except UnicodeDecodeError:
                    pass
            raise DomainError("无法识别 CSV 编码")
        if not filename.lower().endswith(".xlsx"):
            raise DomainError("该来源需要 CSV 或 XLSX 账单")
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                    raise DomainError("Excel 解压后过大")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                rows = []
                for sheet in book:
                    # Some official exports declare A1 even when the XML contains many rows.
                    sheet.reset_dimensions()
                    for row in sheet.iter_rows(values_only=True):
                        if len(rows) >= 25000 or len(row) > 100:
                            raise DomainError("Excel 超过支持的行列范围")
                        rows.append(list(row))
                return rows
            finally:
                book.close()
        except DomainError:
            raise
        except Exception:
            raise DomainError("无法读取 Excel，请上传未加密的原始 XLSX 文件") from None

    @staticmethod
    def text(value):
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return str(value).strip().lstrip("\ufeff")

    @classmethod
    def records(cls, rows, required):
        header = None
        for number, row in enumerate(rows, 1):
            cells = [cls.text(v) for v in row]
            if required.issubset(set(cells)):
                header = cells
                continue
            if header is None or not any(cells):
                continue
            if cells[0].startswith(("---", "合计", "总计", "说明", "注：")):
                continue
            for index, key in enumerate(header):
                if (
                    key in ("交易单号", "订单号")
                    and index < len(row)
                    and isinstance(row[index], (int, float))
                ):
                    raise DomainError(
                        f"第 {number} 行流水号为数值，可能损失精度；请使用原始文本流水号"
                    )
            yield number, {
                k: cells[i] if i < len(cells) else "" for i, k in enumerate(header) if k
            }
        if header is None:
            raise DomainError("未找到该来源账单表头")


class CommonStatementParser(StatementParser):
    provider = ""
    label = ""

    def row(
        self,
        line,
        raw,
        owner,
        *,
        time,
        amount,
        direction,
        description,
        counterparty="",
        method="",
        order_id="",
        status="交易成功",
        category="",
        event="payment",
        extra=None,
    ):
        d = dict(
            line=line,
            raw=raw,
            occurred_at=time,
            amount=0,
            direction=direction,
            note=description,
            counterparty=counterparty,
            method=method,
            order_id=order_id,
            status=status,
            source_category=category,
            error="",
        )
        d.update(extra or {})
        try:
            if self.provider in ("wechat", "alipay_pouch") and not order_id:
                raise DomainError("缺少交易流水号，请上传完整原始账单")
            d["occurred_at"] = Clock.normalize(time)
            d["amount"] = Money.parse(str(amount).replace("¥", "").replace("￥", ""))
            if d["amount"] < 0 or direction not in (
                "收入",
                "支出",
                "不计收支",
                "中性交易",
                "/",
            ):
                raise DomainError("金额或收支方向无效")
            if order_id and (
                re.fullmatch(r"\d+(?:\.\d+)?[eE][+-]?\d+", order_id)
                or order_id.endswith(".0")
            ):
                raise DomainError("流水编号已被转换为数字，请重新导出文本编号")
        except DomainError as exc:
            d["error"] = str(exc)
        identity = (
            [owner, order_id, event]
            if order_id
            else [
                owner,
                d["occurred_at"],
                direction,
                d["amount"],
                d.get("balance"),
                description,
                counterparty,
                d.get("counterparty_account", ""),
            ]
        )
        d["source_key"] = (
            self.provider
            + ":"
            + hashlib.sha256(
                json.dumps(identity, ensure_ascii=False, separators=(",", ":")).encode()
            ).hexdigest()
        )
        return d

    def statement(self, owner, rows, metadata=None):
        if not rows:
            raise DomainError("未解析到交易记录，请检查账单格式")
        totals = defaultdict(lambda: {"count": 0, "amount": 0})
        for row in rows:
            if not row["error"]:
                totals[row["direction"]]["count"] += 1
                totals[row["direction"]]["amount"] += row["amount"]
        return ParsedStatement(
            owner,
            rows,
            dict(
                parser=self.provider,
                provider_label=self.label,
                version="1.0",
                totals=dict(totals),
                warnings=[],
                **(metadata or {}),
            ),
        )


class WechatParser(CommonStatementParser):
    provider, label = "wechat", "微信支付"
    REQUIRED = {
        "交易时间",
        "交易类型",
        "交易对方",
        "商品",
        "收/支",
        "支付方式",
        "当前状态",
        "交易单号",
    }

    def parse(self, content, filename):
        table = SpreadsheetReader.read(content, filename)
        text = "\n".join(SpreadsheetReader.text(r[0]) for r in table if r)
        owner_match = re.search(r"微信昵称：[\[【](.+?)[\]】]", text)
        if not owner_match:
            raise DomainError("缺少微信账单所属昵称")
        owner = owner_match.group(1)
        rows = []
        for line, raw in SpreadsheetReader.records(table, self.REQUIRED):
            direction, status, category = raw["收/支"], raw["当前状态"], raw["交易类型"]
            is_refund = "退款" in category
            method = raw["支付方式"]
            review = ""
            if re.search(r"银行|储蓄卡|信用卡", method):
                review = "银行卡渠道付款或退款，可能同时出现在银行流水中；请核对来源，避免重复入账"
            if direction in ("/", "中性交易", "不计收支"):
                direction = "中性交易"
                review = "微信中性交易，需确认转出和转入账户，不自动计作收支"
            normalized_status = (
                "退款成功"
                if is_refund
                else (
                    "交易成功"
                    if status in ("支付成功", "已转账", "对方已收钱", "已存入零钱")
                    or (direction == "支出" and status.startswith("已退款"))
                    else "交易关闭" if status in ("已关闭", "已全额退还") else status
                )
            )
            rows.append(
                self.row(
                    line,
                    raw,
                    owner,
                    time=raw["交易时间"],
                    amount=raw.get("金额(元)", raw.get("金额（元）", "")),
                    direction=direction,
                    description=" · ".join(
                        v
                        for v in [
                            category,
                            raw["商品"] if raw["商品"] != "/" else "",
                            raw.get("备注", "") if raw.get("备注") != "/" else "",
                        ]
                        if v
                    ),
                    counterparty=raw["交易对方"],
                    method=method,
                    order_id=raw["交易单号"],
                    status=normalized_status,
                    category=category,
                    event="refund" if is_refund else "payment",
                    extra={
                        "review_reason": review,
                        "funding_hint": (
                            "微信"
                            if method == "零钱"
                            or (direction == "收入" and status == "已存入零钱")
                            else method
                        ),
                    },
                )
            )
        result = self.statement(
            owner, rows, {"identity_basis": "微信交易单号；昵称仅用于账户映射"}
        )
        # Provider transaction numbers are globally unique; nicknames can change.
        for row in result.rows:
            if row["order_id"]:
                row["source_key"] = (
                    "wechat:"
                    + hashlib.sha256(
                        (
                            row["order_id"]
                            + (
                                "|refund"
                                if "退款" in row["source_category"]
                                else "|payment"
                            )
                        ).encode()
                    ).hexdigest()
                )
        count = re.search(r"共\s*(\d+)\s*笔记录", text)
        if count and int(count.group(1)) != len(rows):
            raise DomainError("微信声明笔数与解析笔数不一致")
        for direction, count, amount in re.findall(
            r"(收入|支出|中性交易)：(\d+)笔\s*([\d,.]+)元", text
        ):
            actual = result.metadata["totals"].get(direction, {"count": 0, "amount": 0})
            if actual != {"count": int(count), "amount": Money.parse(amount)}:
                raise DomainError("微信" + direction + "汇总与明细不一致")
        return result


class PouchParser(CommonStatementParser):
    provider, label = "alipay_pouch", "支付宝小荷包"
    REQUIRED = {"订单号", "交易时间", "交易说明", "收入金额", "支出金额"}

    def parse(self, content, filename):
        table = SpreadsheetReader.read(content, filename)
        text = "\n".join(SpreadsheetReader.text(r[0]) for r in table if r)
        account = re.search(r"支付宝小荷包账户ID[：:]\s*(\d+)", text)
        name = re.search(r"支付宝小荷包名称[：:]\s*(.+)", text)
        if not account or not name:
            raise DomainError("缺少小荷包账户 ID 或名称")
        owner = account.group(1)
        rows = []
        for line, raw in SpreadsheetReader.records(table, self.REQUIRED):
            income, expense = raw["收入金额"], raw["支出金额"]
            if bool(income) == bool(expense):
                raise DomainError(f"小荷包第 {line} 行需且只能有一个收支金额")
            description = raw["交易说明"]
            review = (
                "小荷包转存或转出，请确认另一侧账户；不自动计作收入或消费"
                if any(
                    t in description
                    for t in ("转入", "转出", "转存", "存入", "取出", "自动攒")
                )
                else ""
            )
            rows.append(
                self.row(
                    line,
                    raw,
                    owner,
                    time=raw["交易时间"],
                    amount=income or expense,
                    direction="收入" if income else "支出",
                    description=description
                    + ((" · " + raw["备注"]) if raw.get("备注") else ""),
                    method="支付宝小荷包(" + name.group(1) + ")",
                    order_id=raw["订单号"],
                    status=(
                        "退款成功" if description.startswith("退款-") else "交易成功"
                    ),
                    event="refund" if description.startswith("退款-") else "payment",
                    extra={
                        "review_reason": review,
                        "operator": raw.get("操作人姓名", ""),
                        "account_number": owner,
                    },
                )
            )
        return self.statement(
            owner,
            rows,
            {
                "account_number": owner,
                "account_name": name.group(1),
                "identity_basis": "小荷包账户 ID + 订单号",
            },
        )


class BankPdfReader:
    @staticmethod
    def open(content):
        import pdfplumber

        try:
            document = pdfplumber.open(io.BytesIO(content))
            if len(document.pages) > 100:
                document.close()
                raise DomainError("PDF 页数超过 100 页，请拆分导出")
            if not document.pages or not document.pages[0].extract_text():
                document.close()
                raise DomainError("PDF 没有可读取文字，请使用银行导出的原始 PDF")
            return document
        except DomainError:
            raise
        except Exception:
            raise DomainError("无法读取 PDF，请使用未加密的银行原始流水") from None


class BankStatementParser(CommonStatementParser):
    def bank_row(
        self,
        line,
        raw,
        owner,
        date,
        amount,
        balance,
        description,
        counterparty,
        counterparty_account="",
        precision="second",
    ):
        signed = Money.parse(amount)
        refund = "退款" in description or "退货" in description
        review = ""
        if (
            "转账" in description
            or "汇款" in description
            or counterparty == raw.get("户名")
        ):
            review = (
                "银行转账或汇款，请确认是否为自有账户往来及另一侧账户，避免重复计收支"
            )
        return self.row(
            line,
            raw,
            owner,
            time=date,
            amount=format(Decimal(abs(signed)) / 100, ".2f"),
            direction="收入" if signed >= 0 else "支出",
            description=description,
            counterparty=counterparty,
            method=self.label + "(" + owner[-4:] + ")",
            status="退款成功" if refund else "交易成功",
            category=description,
            extra={
                "balance": Money.parse(balance),
                "signed_amount": signed,
                "account_number": owner,
                "counterparty_account": counterparty_account,
                "time_precision": precision,
                "review_reason": review,
            },
        )

    def finish(self, owner, rows, extra=None):
        ordered = rows if self.provider == "cmb" else list(reversed(rows))
        for prev, current in zip(ordered, ordered[1:]):
            if prev["balance"] + current["signed_amount"] != current["balance"]:
                raise DomainError("银行流水余额衔接不一致，请核对是否漏行或格式变化")
        if not ordered:
            raise DomainError("银行流水没有可导入记录")
        identities = [r["source_key"] for r in rows]
        if len(set(identities)) != len(identities):
            raise DomainError("银行流水存在完全相同记录，缺少流水号，需人工核对后导入")
        result = self.statement(
            owner,
            rows,
            dict(
                account_number=owner,
                account_name=self.label,
                identity_basis="卡号、日期/时间、方向、金额、余额和对方信息的组合；原文件未提供交易流水号",
                opening_balance=ordered[0]["balance"] - ordered[0]["signed_amount"],
                closing_balance=ordered[-1]["balance"],
                balance_chain_checked=True,
                **(extra or {}),
            ),
        )
        result.metadata["warnings"].append(
            "银行文件没有交易流水号，使用账户、日期/时间、金额、余额及对方信息组合去重；跨平台付款和转账需核对"
        )
        if self.provider == "cmb":
            result.metadata["warnings"].append(
                "招行只提供记账日期；统一时间中的 00:00:00 是日期占位，不代表真实交易时间"
            )
        return result


class CmbParser(BankStatementParser):
    provider, label = "cmb", "招商银行"
    HEADERS = ["记账日期", "货币", "交易金额", "联机余额", "交易摘要", "对手信息"]

    def parse(self, content, filename):
        rows = []
        with BankPdfReader.open(content) as pdf:
            text = pdf.pages[0].extract_text()
            if "招商银行交易流水" not in text:
                raise DomainError("不是支持的招商银行交易流水")
            account = re.search(r"账号[：:]\s*(\d+)", text)
            person = re.search(r"户\s*名[：:]\s*(\S+)", text)
            if not account:
                raise DomainError("未找到招行账号")
            owner = account.group(1)
            for page_no, page in enumerate(pdf.pages, 1):
                words = page.extract_words()
                header = {w["text"]: w for w in words if w["text"] in self.HEADERS}
                if len(header) != len(self.HEADERS):
                    raise DomainError(f"招行第 {page_no} 页缺少完整表头")
                starts = [header[h]["x0"] - 2 for h in self.HEADERS] + [page.width]
                anchors = [
                    w
                    for w in words
                    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", w["text"])
                    and starts[0] <= w["x0"] < starts[1]
                    and w["top"] > header["记账日期"]["top"]
                ]
                for i, anchor in enumerate(anchors):
                    top = (
                        (anchors[i - 1]["top"] + anchor["top"]) / 2
                        if i
                        else anchor["top"] - 14
                    )
                    bottom = (
                        (anchors[i + 1]["top"] + anchor["top"]) / 2
                        if i + 1 < len(anchors)
                        else anchor["top"] + 18
                    )
                    cells = []
                    for left, right in zip(starts, starts[1:]):
                        selected = sorted(
                            (
                                w
                                for w in words
                                if left <= w["x0"] < right and top <= w["top"] < bottom
                            ),
                            key=lambda w: (round(w["top"], 1), w["x0"]),
                        )
                        cells.append(
                            (" " if left == starts[-2] else "").join(
                                w["text"] for w in selected
                            )
                        )
                    if cells[1] != "CNY":
                        raise DomainError("当前只支持人民币银行流水")
                    cp_match = re.search(r"(?<=\s)(\d{5,})$", cells[5])
                    cp_number = cp_match.group(1) if cp_match else ""
                    cp_name = (
                        cells[5][: -len(cp_number)] if cp_number else cells[5]
                    ).replace(" ", "")
                    raw = dict(
                        zip(self.HEADERS, cells),
                        页码=page_no,
                        户名=person.group(1) if person else "",
                    )
                    rows.append(
                        self.bank_row(
                            len(rows) + 1,
                            raw,
                            owner,
                            cells[0] + " 00:00:00",
                            cells[2],
                            cells[3],
                            cells[4],
                            cp_name,
                            cp_number,
                            "date",
                        )
                    )
            return self.finish(owner, rows, {"pages": len(pdf.pages)})


class BocParser(BankStatementParser):
    provider, label = "boc", "中国银行"

    def parse(self, content, filename):
        rows = []
        with BankPdfReader.open(content) as pdf:
            text = pdf.pages[0].extract_text()
            if "中国银行交易流水明细清单" not in text:
                raise DomainError("不是支持的中国银行交易流水")
            account = re.search(r"借记卡号[：:]\s*(\d+)", text)
            person = re.search(r"客户姓名[：:]\s*(\S+)", text)
            if not account:
                raise DomainError("未找到中国银行借记卡号")
            owner = account.group(1)
            for page_no, page in enumerate(pdf.pages, 1):
                found = False
                for table in page.extract_tables():
                    if not table or not table[0] or table[0][0] != "记账日期":
                        continue
                    found = True
                    header = table[0]
                    for values in table[1:]:
                        if not any(values):
                            continue
                        raw = {
                            key: (value or "").replace("\n", "").strip()
                            for key, value in zip(header, values)
                        }
                        if not re.fullmatch(
                            r"\d{4}-\d{2}-\d{2}", raw.get("记账日期", "")
                        ):
                            raise DomainError(
                                f"中国银行第 {page_no} 页出现无法识别的数据行"
                            )
                        if raw.get("币别") != "人民币":
                            raise DomainError("当前只支持人民币银行流水")
                        raw.update(页码=page_no, 户名=person.group(1) if person else "")
                        clean = lambda value: (
                            "" if re.fullmatch("-+", value or "") else value
                        )
                        rows.append(
                            self.bank_row(
                                len(rows) + 1,
                                raw,
                                owner,
                                raw["记账日期"] + " " + raw["记账时间"],
                                raw["金额"],
                                raw["余额"],
                                raw["交易名称"]
                                + (
                                    (" · " + clean(raw["附言"]))
                                    if clean(raw.get("附言", ""))
                                    else ""
                                ),
                                clean(raw["对方账户名"]),
                                clean(raw["对方卡号/账号"]),
                            )
                        )
                if not found:
                    raise DomainError(f"中国银行第 {page_no} 页没有可识别表格")
            count = re.search(r"行数[：:]\s*(\d+)", text)
            if count and int(count.group(1)) != len(rows):
                raise DomainError("中国银行声明行数与解析行数不一致")
            for name, direction in [("借方发生数", "支出"), ("贷方发生数", "收入")]:
                declared = re.search(name + r"[：:]\s*([\d,.]+)", text)
                if declared and Money.parse(declared.group(1)) != sum(
                    r["amount"] for r in rows if r["direction"] == direction
                ):
                    raise DomainError("中国银行" + name + "与明细不一致")
            return self.finish(owner, rows, {"pages": len(pdf.pages)})
